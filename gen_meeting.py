"""
회의록 Word 문서 생성기 (OK금융그룹 양식)
사용법: python gen_meeting.py meeting_data.json [output.docx]
JSON 스키마는 sample_meeting.json 참고
"""
import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml

# ── 스타일 상수 ──
FONT = "맑은 고딕"
DARK = "1F4E79"
LIGHT = "D6E4F0"
GRAY = "666666"
SZ = 21  # 10.5pt (half-points)
TABLE_W = 9072  # 전체 테이블 너비 (twips)
COL_LABEL = 2100
COL_CONTENT = TABLE_W - COL_LABEL  # 6972


# ── 유틸리티 ──
def sf(run, sz=SZ, bold=False, color=None):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    rPr = run._element.rPr
    rPr.append(parse_xml(f'<w:sz {nsdecls("w")} w:val="{sz}"/>'))
    rPr.append(parse_xml(f'<w:szCs {nsdecls("w")} w:val="{sz}"/>'))
    run.bold = bold
    if color:
        rPr.append(parse_xml(f'<w:color {nsdecls("w")} w:val="{color}"/>'))


def shade(cell, c):
    cell._element.get_or_add_tcPr().append(
        parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{c}"/>')
    )


def cw(cell, w):
    cell._element.get_or_add_tcPr().append(
        parse_xml(f'<w:tcW {nsdecls("w")} w:w="{w}" w:type="dxa"/>')
    )


def valign(cell):
    cell._element.get_or_add_tcPr().append(
        parse_xml(f'<w:vAlign {nsdecls("w")} w:val="center"/>')
    )


def compact(p, before=1, after=1):
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)


def label_cell(cell, text, width=1553):
    cw(cell, width)
    shade(cell, LIGHT)
    valign(cell)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    compact(p, 4, 4)
    sf(p.add_run(text), SZ, bold=True)


def value_cell(cell, text, width=2591):
    cw(cell, width)
    valign(cell)
    p = cell.paragraphs[0]
    compact(p, 4, 4)
    sf(p.add_run(text), SZ)


def header_2col(table, row=0):
    for ci, (txt, w) in enumerate([("항목", COL_LABEL), ("내용", COL_CONTENT)]):
        c = table.cell(row, ci)
        cw(c, w)
        shade(c, DARK)
        valign(c)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        compact(p, 4, 4)
        sf(p.add_run(txt), SZ, bold=True, color="FFFFFF")


def body_2col(table, row, cat, detail):
    c0 = table.cell(row, 0)
    cw(c0, COL_LABEL)
    shade(c0, LIGHT)
    valign(c0)
    p = c0.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    compact(p, 6, 6)
    parts = cat.split("\n")
    for j, line in enumerate(parts):
        sf(p.add_run(line), SZ, bold=True)
        if j < len(parts) - 1:
            sf(p.add_run("\n"), SZ)

    c1 = table.cell(row, 1)
    cw(c1, COL_CONTENT)
    valign(c1)
    first_p = c1.paragraphs[0]
    lines = detail.split("\n")
    cur = first_p
    is_first = True
    for j, line in enumerate(lines):
        if re.match(r"^\d+\.", line) and not is_first:
            cur = c1.add_paragraph()
            compact(cur, 4, 0)
        elif is_first:
            compact(cur, 4, 0)
        sf(cur.add_run(line), SZ)
        if j < len(lines) - 1 and not re.match(r"^\d+\.", lines[j + 1]):
            sf(cur.add_run("\n"), SZ)
        is_first = False


def header_4col(table, cols):
    for ci, (txt, w) in enumerate(cols):
        c = table.cell(0, ci)
        cw(c, w)
        shade(c, DARK)
        valign(c)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        compact(p, 4, 4)
        sf(p.add_run(txt), SZ, bold=True, color="FFFFFF")


def body_4col(table, row, values, widths):
    for ci, (val, w) in enumerate(zip(values, widths)):
        c = table.cell(row, ci)
        cw(c, w)
        if ci == 0:
            shade(c, LIGHT)
        valign(c)
        p = c.paragraphs[0]
        compact(p, 3, 3)
        vlines = val.split("\n")
        for j, line in enumerate(vlines):
            sf(p.add_run(line), SZ, bold=(ci == 0))
            if j < len(vlines) - 1:
                sf(p.add_run("\n"), SZ)


def section_header(cell, text):
    cw(cell, TABLE_W)
    shade(cell, DARK)
    valign(cell)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    compact(p, 4, 4)
    sf(p.add_run(text), SZ, bold=True, color="FFFFFF")


def section_body(cell, text):
    cw(cell, TABLE_W)
    p = cell.paragraphs[0]
    compact(p, 4, 4)
    slines = text.strip().split("\n")
    for j, line in enumerate(slines):
        sf(p.add_run(line), SZ)
        if j < len(slines) - 1:
            sf(p.add_run("\n"), SZ)


# ── 메인 생성 함수 ──
def generate(data: dict, output_path: str):
    doc = Document()

    for s in doc.sections:
        s.page_width = 7560310
        s.page_height = 10692130
        s.top_margin = 360000
        s.bottom_margin = 360000
        s.left_margin = 914400
        s.right_margin = 914400

    info = data["기본정보"]

    # ── 제목 ──
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(4)
    sf(title.add_run("회의록"), 60, bold=True, color=DARK)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    sub.paragraph_format.space_before = Pt(0)
    sub.paragraph_format.space_after = Pt(6)
    sf(sub.add_run(f'작성일 : {info["작성일"]}  /  작성자 : {info["작성자"]}'), 16)

    # ── 기본정보 테이블 ──
    t0 = doc.add_table(rows=4, cols=4)
    t0.alignment = WD_TABLE_ALIGNMENT.CENTER
    t0.style = "Table Grid"

    label_cell(t0.cell(0, 0), "회의 일시")
    value_cell(t0.cell(0, 1), info["회의일시"])
    label_cell(t0.cell(0, 2), "회의 장소", 1481)
    value_cell(t0.cell(0, 3), info.get("회의장소", "-"), 3447)

    label_cell(t0.cell(1, 0), "주관 부서")
    value_cell(t0.cell(1, 1), info.get("주관부서", "-"))
    label_cell(t0.cell(1, 2), "참석 부서", 1481)
    value_cell(t0.cell(1, 3), info.get("참석부서", "-"), 3447)

    label_cell(t0.cell(2, 0), "참 석 자")
    t0.cell(2, 1).merge(t0.cell(2, 3))
    c = t0.cell(2, 1)
    valign(c)
    p = c.paragraphs[0]
    compact(p, 4, 4)
    sf(p.add_run(info["참석자"]), SZ)

    label_cell(t0.cell(3, 0), "회의 안건")

    t0.cell(3, 1).merge(t0.cell(3, 3))
    c = t0.cell(3, 1)
    valign(c)
    p = c.paragraphs[0]
    compact(p, 4, 4)
    sf(p.add_run(info["회의안건"]), SZ)

    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(2)
    sp.paragraph_format.space_after = Pt(2)

    # ── 회의내용 테이블 ──
    content = data["회의내용"]
    t1 = doc.add_table(rows=1 + len(content), cols=2)
    t1.alignment = WD_TABLE_ALIGNMENT.CENTER
    t1.style = "Table Grid"
    header_2col(t1)
    for i, item in enumerate(content):
        body_2col(t1, i + 1, item["항목"], item["내용"])

    # ── 페이지 나누기 ──
    pb = doc.add_paragraph()
    pb.paragraph_format.space_before = Pt(0)
    pb.paragraph_format.space_after = Pt(0)
    pb.paragraph_format.line_spacing = Pt(1)
    pb.add_run()._element.append(
        parse_xml(f'<w:br {nsdecls("w")} w:type="page"/>')
    )

    # ── 결정사항 테이블 ──
    decisions = data["결정사항"]
    col_defs = [("No", 450), ("결정사항", 5322), ("담당자", 2100), ("비고", 1200)]
    widths = [c[1] for c in col_defs]

    t_act = doc.add_table(rows=1 + len(decisions), cols=4)
    t_act.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_act.style = "Table Grid"
    header_4col(t_act, col_defs)
    for i, item in enumerate(decisions):
        vals = [str(i + 1), item["내용"], item["담당자"], item.get("비고", "-")]
        body_4col(t_act, i + 1, vals, widths)

    # ── 기타 ──
    if "기타" in data and data["기타"]:
        sp = doc.add_paragraph()
        sp.paragraph_format.space_before = Pt(2)
        sp.paragraph_format.space_after = Pt(2)

        t2 = doc.add_table(rows=2, cols=1)
        t2.alignment = WD_TABLE_ALIGNMENT.CENTER
        t2.style = "Table Grid"
        section_header(t2.cell(0, 0), "기타")
        section_body(t2.cell(1, 0), "\n".join(f"- {x}" for x in data["기타"]))

    # ── 하단 ──
    ft = doc.add_paragraph()
    ft.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ft.paragraph_format.space_before = Pt(10)
    ft.paragraph_format.space_after = Pt(0)
    sf(ft.add_run("- 이상 끝 -"), SZ, color=GRAY)

    doc.save(output_path)
    print(f"SAVED: {output_path}")


# ── CLI ──
if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python gen_meeting.py <meeting_data.json> [output.docx]")
        sys.exit(1)

    json_path = sys.argv[1]
    with open(json_path, "r", encoding="utf-8") as f:
        data = json.load(f)

    if len(sys.argv) >= 3:
        out = sys.argv[2]
    else:
        title = data["기본정보"].get("회의안건", "회의록")[:20].replace(" ", "_")
        date = data["기본정보"]["작성일"].replace(". ", "").replace(".", "")
        out = f"회의록_{date}.docx"

    generate(data, out)
